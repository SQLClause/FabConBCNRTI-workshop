#!/usr/bin/env python3
"""
Regenerate the PDF and PPTX for the Claude Slides deck from fabric-iq-workshop.html.

  pip install python-pptx pymupdf
  python3 build.py

- PDF:  printed with headless Google Chrome (1920x1080 pages, web fonts fetched from Google Fonts).
- PPTX: one full-bleed picture per slide (rendered from the PDF) plus the slide's speaker notes.
        It is NOT an editable deck -- edit the HTML (or the Claude Slides artifact) and rebuild.
"""
import os, re, sys, html, subprocess, shutil, tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
HTML = os.path.join(HERE, "fabric-iq-workshop.html")
DIST = os.path.normpath(os.path.join(HERE, "..", "dist"))
PDF = os.path.join(DIST, "pdf", "fabric-iq-workshop.pdf")
PPTX = os.path.join(DIST, "pptx", "fabric-iq-workshop.pptx")
CHROME = next((c for c in [
    os.environ.get("CHROME"),
    "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
    shutil.which("google-chrome"), shutil.which("chromium"), shutil.which("chrome"),
] if c and os.path.exists(c)), None)
if not CHROME:
    sys.exit("Google Chrome not found; set CHROME=/path/to/chrome")

import pymupdf
from pptx import Presentation
from pptx.util import Inches

os.makedirs(os.path.dirname(PDF), exist_ok=True); os.makedirs(os.path.dirname(PPTX), exist_ok=True)
subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                "--virtual-time-budget=15000", "--run-all-compositor-stages-before-draw",
                f"--print-to-pdf={PDF}", f"file://{HTML}"], check=True, capture_output=True, timeout=300)

src = open(HTML, encoding="utf-8").read()
notes = [html.unescape(re.search(r"<aside>(.*?)</aside>", s, re.S).group(1)).strip()
         if re.search(r"<aside>(.*?)</aside>", s, re.S) else ""
         for s in re.findall(r"<section\b.*?</section>", src, re.S)]

pdf = pymupdf.open(PDF)
assert len(pdf) == len(notes), f"{len(pdf)} PDF pages vs {len(notes)} sections"
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
with tempfile.TemporaryDirectory() as tmp:
    for n, page in enumerate(pdf):
        png = os.path.join(tmp, f"{n:03d}.png")
        page.get_pixmap(matrix=pymupdf.Matrix(4 / 3, 4 / 3), alpha=False).save(png)  # 1920x1080
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(png, 0, 0, prs.slide_width, prs.slide_height)
        if notes[n]:
            slide.notes_slide.notes_text_frame.text = notes[n]
prs.save(PPTX)
print(f"{len(pdf)} slides -> {PDF}\n{len(pdf)} slides -> {PPTX}")
